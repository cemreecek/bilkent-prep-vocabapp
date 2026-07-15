import docx

def inspect_docx(docx_path):
    print(f"Inspecting {docx_path}...")
    doc = docx.Document(docx_path)
    
    print("--- Paragraphs ---")
    for i, p in enumerate(doc.paragraphs[:50]): # print first 50 paragraphs
        text = p.text.strip()
        if text:
            print(f"P{i}: {text}")
            
    print("--- Tables ---")
    for i, table in enumerate(doc.tables[:2]):
        print(f"Table {i}: {len(table.rows)} rows")
        for j, row in enumerate(table.rows[:2]):
            row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"  Row {j}: {row_data}")

if __name__ == "__main__":
    test_docx = r"d:\Bilkent\BilkentApp\BilkentApp\Prefac\WORDLIST SETS & VOCABULARY STRAND\PFC WORD LIST SETS SPRING PRACTICE MATERIALS (Updated June 2025)\PFC Spring Semester Wordlist Practice Set 1.docx"
    try:
        inspect_docx(test_docx)
    except Exception as e:
        print(f"Error: {e}")
