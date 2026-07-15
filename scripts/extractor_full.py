import os
import glob
import json
import docx
import win32com.client

def get_word():
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    return word

def doc_to_docx(word, doc_path):
    docx_path = doc_path + "x"
    if not os.path.exists(docx_path):
        print(f"Converting {doc_path}...")
        try:
            doc = word.Documents.Open(doc_path)
            doc.SaveAs2(docx_path, FileFormat=16)
            doc.Close()
        except Exception as e:
            print(f"Failed to convert {doc_path}: {e}")
            return None
    return docx_path

def parse_wordlist(docx_path):
    doc = docx.Document(docx_path)
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    flashcards = []
    for i, table in enumerate(doc.tables):
        if i >= len(days):
            break
        day = days[i]
        for row in table.rows:
            word_text = row.cells[0].text.strip()
            if word_text and word_text.lower() not in [d.lower() for d in days] and not word_text.startswith("SET "):
                flashcards.append({
                    "id": word_text.lower().replace('/', '_'),
                    "day": day,
                    "word": word_text,
                    "definition": "Manual touch-up needed" # placeholder as tables only have words
                })
    return flashcards

def parse_practice(docx_path):
    doc = docx.Document(docx_path)
    practices = []
    
    # Simple heuristic parser
    current_practice = None
    manual_touch_ups = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
            
        text_upper = text.upper()
        if "PART" in text_upper and "MATCH" in text_upper:
            current_practice = {"id": "practice-match", "type": "matching", "title": text, "content": "Manual extraction needed"}
            practices.append(current_practice)
        elif "PART" in text_upper and "CHOOSE" in text_upper:
            current_practice = {"id": "practice-mc", "type": "multiple-choice", "title": text, "content": "Manual extraction needed"}
            practices.append(current_practice)
        elif "PART" in text_upper and "FILL" in text_upper:
            current_practice = {"id": "practice-gapfill", "type": "gap-fill", "title": text, "content": "Manual extraction needed"}
            practices.append(current_practice)
        elif "CROSS" in text_upper and "WORD" in text_upper:
            current_practice = None # Skip crossword
            
    return practices, manual_touch_ups

def main():
    prefac_dir = r"d:\Bilkent\BilkentApp\BilkentApp\Prefac\WORDLIST SETS & VOCABULARY STRAND"
    wordlist_dir = os.path.join(prefac_dir, "PFC LEVEL SPRING WORDLIST SETS (Updated June 2025)")
    practice_dir = os.path.join(prefac_dir, "PFC WORD LIST SETS SPRING PRACTICE MATERIALS (Updated June 2025)")
    
    word = get_word()
    sets_data = []
    touch_ups = []
    
    try:
        # We will process Set 1 to Set 8
        for i in range(1, 9):
            wordlist_doc = os.path.join(wordlist_dir, f"PFC Spring Semester Wordlist- Set {i}.doc")
            practice_docx = os.path.join(practice_dir, f"PFC Spring Semester Wordlist Practice Set {i}.docx")
            if not os.path.exists(practice_docx):
                practice_docx = os.path.join(practice_dir, f"PFC Spring Semester Wordlist Practice Set {i}.doc")
            
            if not os.path.exists(wordlist_doc):
                continue
                
            print(f"Processing Set {i}...")
            wl_docx = doc_to_docx(word, wordlist_doc)
            pr_docx = practice_docx
            if practice_docx.endswith(".doc"):
                pr_docx = doc_to_docx(word, practice_docx)
                
            flashcards = parse_wordlist(wl_docx) if wl_docx else []
            practices = []
            if pr_docx and os.path.exists(pr_docx):
                pr, tu = parse_practice(pr_docx)
                practices = pr
                if tu:
                    touch_ups.append(f"Set {i}: {tu}")
            
            sets_data.append({
                "id": f"prefac-set-{i}",
                "level": "prefac",
                "title": f"PFC Spring Semester Set {i}",
                "flashcards": flashcards,
                "practices": practices
            })
            touch_ups.append(f"Set {i}: Flashcard definitions are missing (only words were in tables). Practices require manual structural parsing.")
    finally:
        word.Quit()
        
    out_file = r"d:\Bilkent\BilkentApp\BilkentApp\app\public\data\vocabulary_full.json"
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump({"sets": sets_data}, f, indent=2)
        
    with open("touch_ups.txt", 'w', encoding='utf-8') as f:
        for t in touch_ups:
            f.write(t + "\n")
            
    print("Done. See touch_ups.txt")

if __name__ == "__main__":
    main()
