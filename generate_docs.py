import os
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from docx.shared import Pt as DocxPt
from pptx.enum.text import PP_ALIGN

def create_guidelines():
    doc = Document()
    
    title = doc.add_heading('Bilkent Prep Vocab App: User Guidelines', 0)
    title.alignment = 1 # Center
    
    doc.add_paragraph('Welcome to the Bilkent Prep Vocab App! This platform is designed to provide a seamless, gamified vocabulary learning experience for students, while offering powerful analytics and management tools for instructors and administrators.')
    
    # STUDENT SECTION
    doc.add_heading('1. Student Experience', level=1)
    doc.add_paragraph('Students log in to a personalized dashboard tailored to their assigned level.')
    
    doc.add_heading('Key Features:', level=2)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Curriculum Modules: ').bold = True
    p.add_run('Access is restricted to current and previous levels (e.g., Elementary through Intermediate). Future levels remain locked.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Streaks & Gamification: ').bold = True
    p.add_run('Daily logins and practice sessions increase the student\'s streak. Total error points (mistakes) are also tracked to highlight areas for improvement.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Practice Sessions: ').bold = True
    p.add_run('All vocabulary practice is grouped by instructions (e.g., matching, cloze). Questions and options are fully randomized on every attempt.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Quick Resume: ').bold = True
    p.add_run('The "Continue Learning" button instantly resumes practice in their currently assigned level.')

    # TEACHER SECTION
    doc.add_heading('2. Instructor Portal (Teacher View)', level=1)
    doc.add_paragraph('Instructors have access to live analytics for their assigned classrooms.')
    
    doc.add_heading('Key Features:', level=2)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Live Class Analytics: ').bold = True
    p.add_run('View Total Students, Class Average Streak, and overall Engagement metrics.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Roster & Progress: ').bold = True
    p.add_run('Monitor individual student metrics, including last login times, time spent studying, and accumulated error debt.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('CSV Export: ').bold = True
    p.add_run('A single click generates a downloadable CSV of the class roster containing all vital student metrics for offline tracking.')

    # ADMIN SECTION
    doc.add_heading('3. Admin "God View"', level=1)
    doc.add_paragraph('System overseers manage the entire platform, creating classrooms and assigning roles.')
    
    doc.add_heading('Key Features:', level=2)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Classroom Management: ').bold = True
    p.add_run('Create sections, assign teachers, and view global institution metrics (total classrooms, total active users, global error debt).')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('User Management: ').bold = True
    p.add_run('View the complete user directory. Easily promote students to teachers or admins using the "Change Role / Edit" interface.')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('Student Assigner Modal: ').bold = True
    p.add_run('Add students to sections manually or search through unassigned students by name or email.')

    doc.add_paragraph('\nFor any technical issues, please contact the IT Administrator.')
    
    doc.save('Bilkent_Prep_User_Guidelines.docx')
    print("Created Bilkent_Prep_User_Guidelines.docx")

def create_pitch():
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Bilkent Prep Vocab App"
    subtitle.text = "Revolutionizing Vocabulary Practice\nPresented by Cemre Ecek"
    
    # Slide 2: The Problem
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "The Challenge"
    content = slide.placeholders[1].text_frame
    content.text = "Current vocabulary learning faces multiple hurdles:"
    p = content.add_paragraph()
    p.text = "• Manual tracking of student progress"
    p = content.add_paragraph()
    p.text = "• Lack of engaging, gamified elements for students"
    p = content.add_paragraph()
    p.text = "• Disconnected systems between admins, teachers, and students"

    # Slide 3: The Solution
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "The Solution"
    content = slide.placeholders[1].text_frame
    content.text = "A centralized, responsive, and data-driven platform designed specifically for the Bilkent Prep curriculum."
    p = content.add_paragraph()
    p.text = "• Role-based access (Student, Teacher, Admin)"
    p = content.add_paragraph()
    p.text = "• Automated randomization and dynamic question delivery"
    p = content.add_paragraph()
    p.text = "• Real-time analytics and tracking"

    # Slide 4: Student Value
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "For Students: Gamified Learning"
    content = slide.placeholders[1].text_frame
    content.text = "Driving engagement through active recall:"
    p = content.add_paragraph()
    p.text = "• Level-based progression (locked vs. accessible modules)"
    p = content.add_paragraph()
    p.text = "• Daily Streaks to encourage consistent practice"
    p = content.add_paragraph()
    p.text = "• Immediate feedback on vocabulary sets"
    p = content.add_paragraph()
    p.text = "• Seamless mobile and desktop experience"

    # Slide 5: Teacher Value
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "For Teachers: Actionable Insights"
    content = slide.placeholders[1].text_frame
    content.text = "Empowering instructors with real-time data:"
    p = content.add_paragraph()
    p.text = "• Live dashboards showing class streaks and engagement"
    p = content.add_paragraph()
    p.text = "• Individual student tracking (time spent, error debt)"
    p = content.add_paragraph()
    p.text = "• 1-Click CSV Roster Exports for gradebook integration"

    # Slide 6: Admin Value
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "For Admins: 'God View'"
    content = slide.placeholders[1].text_frame
    content.text = "Complete control over the institutional ecosystem:"
    p = content.add_paragraph()
    p.text = "• Searchable, unified User Management"
    p = content.add_paragraph()
    p.text = "• Dynamic classroom creation and teacher assignments"
    p = content.add_paragraph()
    p.text = "• High-level metrics tracking institutional 'Error Debt'"

    # Slide 7: Tech Stack
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Technical Highlights"
    content = slide.placeholders[1].text_frame
    p = content.add_paragraph()
    p.text = "• Next.js 16 (React) for a lightning-fast UI"
    p = content.add_paragraph()
    p.text = "• Tailwind CSS for a premium, responsive design"
    p = content.add_paragraph()
    p.text = "• PostgreSQL Database with Prisma ORM"
    p = content.add_paragraph()
    p.text = "• Secure Authentication via NextAuth"
    p = content.add_paragraph()
    p.text = "• Hosted globally on Vercel"

    # Slide 8: Next Steps
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Next Steps"
    content = slide.placeholders[1].text_frame
    content.text = "1. Connect custom .edu.tr domain (pending GitHub approval)"
    p = content.add_paragraph()
    p.text = "2. Finalize AI Content Parser for automated vocab generation"
    p = content.add_paragraph()
    p.text = "3. Pilot launch with first Bilkent Prep sections"
    p = content.add_paragraph()
    p.text = "\nThank you! Questions?"

    prs.save('Bilkent_Prep_Pitch_Presentation.pptx')
    print("Created Bilkent_Prep_Pitch_Presentation.pptx")

if __name__ == "__main__":
    create_guidelines()
    create_pitch()
